"""Rewrite repeated formulas in an xlsx as shared formulas (horizontal runs) to shrink the file."""
import re
import sys
import zipfile
import html
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
zin = zipfile.ZipFile(src)
files = {n: zin.read(n) for n in zin.namelist()}

for idx, ws in enumerate(wb.worksheets, start=1):
    name = f"xl/worksheets/sheet{idx}.xml"
    xml = files[name].decode()
    si = 0
    edits = {}  # coord -> ("master", ref, si) | ("child", si)
    for row in ws.iter_rows():
        cells = [c for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        i = 0
        while i < len(cells):
            m = cells[i]
            j = i + 1
            while (j < len(cells) and cells[j].column == cells[j - 1].column + 1
                   and Translator(m.value, origin=m.coordinate).translate_formula(cells[j].coordinate) == cells[j].value):
                j += 1
            if j - i >= 2:
                ref = f"{m.coordinate}:{cells[j - 1].coordinate}"
                edits[m.coordinate] = ("master", ref, si)
                for c in cells[i + 1:j]:
                    edits[c.coordinate] = ("child", si)
                si += 1
            i = j

    def repl(mo):
        coord = mo.group(1)
        e = edits.get(coord)
        if not e:
            return mo.group(0)
        body = mo.group(0)
        if e[0] == "master":
            return re.sub(r"<f>", f'<f t="shared" ref="{e[1]}" si="{e[2]}">', body, count=1)
        return re.sub(r"<f>.*?</f>", f'<f t="shared" si="{e[1]}"/>', body, count=1, flags=re.S)

    xml = re.sub(r'<c r="([A-Z]+\d+)"[^>]*>.*?</c>', repl, xml, flags=re.S)
    xml = xml.replace("<v></v>", "").replace("<v />", "")
    files[name] = xml.encode()

ent = re.compile(r"&#(\d+);")
for n in list(files):
    if n.endswith(".xml"):
        files[n] = ent.sub(lambda m: chr(int(m.group(1))) if int(m.group(1)) > 127 else m.group(0), files[n].decode()).encode()

with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for n, d in files.items():
        z.writestr(n, d)

# verify: openpyxl expands shared formulas on load
w2 = openpyxl.load_workbook(dst)
bad = 0
for a, b in zip(wb.worksheets, w2.worksheets):
    for row in a.iter_rows():
        for c in row:
            if c.value != b[c.coordinate].value:
                bad += 1
                if bad < 5:
                    print("MISMATCH", a.title, c.coordinate, c.value, b[c.coordinate].value)
print("mismatches:", bad)
