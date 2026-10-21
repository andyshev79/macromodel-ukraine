from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = "macromodel_frame.xlsx"

B = Font(bold=True)
TITLE = Font(bold=True, size=14)
INPUT_FONT = Font(color="1F4E9E")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
HEAD_FILL = PatternFill("solid", fgColor="E8EAED")
SEC_FILL = PatternFill("solid", fgColor="D9E7FD")
WRAP = Alignment(wrap_text=True, vertical="top")

YEARS = list(range(2015, 2029))
QUARTERS = [f"{y}Q{q}" for y in YEARS for q in range(1, 5)]

wb = Workbook()


def head(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = B
        c.fill = HEAD_FILL


def inp(c):
    c.font = INPUT_FONT
    c.fill = INPUT_FILL


def widths(ws, w):
    for col, width in w.items():
        ws.column_dimensions[col].width = width


# ---------------- README ----------------
ws = wb.active
ws.title = "README"
rows = [
    ("Макромодель Украины: каркас", None),
    ("Прогноз ключевых показателей на 2–3 года. Квартальная частота, годовая агрегация. Подконтрольная территория.", None),
    (None, None),
    ("Легенда", None),
    ("Ввод", "Синий шрифт на светло-жёлтом фоне: допущение или данные, которые можно менять"),
    ("Формула", "Чёрный шрифт: считается из других ячеек, руками не править"),
    ("Код переменной", "Короткий латинский код (GDP_REAL). Одинаков на всех листах, в Python и в словаре DICT"),
    (None, None),
    ("Листы", None),
    ("README", "Это описание"),
    ("DICT", "Словарь переменных: название, единицы, блок, источник, статус данных"),
    ("SCENARIOS", "Активный сценарий, последний фактический квартал, сценарные допущения по годам"),
    ("ANCHORS", "Реперные переменные по кварталам 2015Q1–2028Q4. Единственный канал связи между блоками"),
    ("ANNUAL", "Годовая агрегация реперов (сумма, среднее или конец периода по колонке «Агрегация»)"),
    ("CHECKS", "Проверки сверху вниз / снизу вверх, мощности (люди, энергия, логистика, капитал) и разрыв финансирования"),
    ("COEFS", "Оценённые коэффициенты уравнений. Заполняет Python, читают листы блоков"),
    ("B1_REAL … B8_CAPACITY", "Листы блоков: уравнения и расчёты. Пишут результат в ANCHORS"),
    ("SOURCES", "Источники данных и ссылки"),
    (None, None),
    ("Правила", None),
    ("1", "Блок читает только ANCHORS, SCENARIOS и COEFS и пишет только в свои строки ANCHORS"),
    ("2", "Пока блок не построен, его строки в ANCHORS заполняются фактом и внешним прогнозом (НБУ, МВФ) как временным допущением"),
    ("3", "Каждое уравнение имеет разложение по вкладам факторов на своём листе блока"),
    ("4", "Python загружает сырые данные (листы RAW_*), оценивает уравнения (COEFS) и сводит итерации; формулы прогноза остаются здесь"),
    (None, None),
    ("Статус блоков", None),
]
for r, (a, b) in enumerate(rows, start=1):
    ws.cell(row=r, column=1, value=a)
    ws.cell(row=r, column=2, value=b)
ws["A1"].font = TITLE
for r in (4, 9, 20, 26):
    ws.cell(row=r, column=1).font = B
    ws.cell(row=r, column=1).fill = SEC_FILL
    ws.cell(row=r, column=2).fill = SEC_FILL
inp(ws["A5"])
head(ws, 27, ["Блок", "Лист", "Статус"])
blocks = [
    ("0. Сценарии и внешняя среда", "SCENARIOS"),
    ("1. Реальный сектор", "B1_REAL"),
    ("2. Спрос (ВВП по расходам)", "B2_DEMAND"),
    ("3. Цены", "B3_PRICES"),
    ("4. Рынок труда", "B4_LABOUR"),
    ("5. Внешний сектор", "B5_EXTERNAL"),
    ("6. Бюджет и долг", "B6_FISCAL"),
    ("7. Монетарный блок", "B7_MONETARY"),
    ("8. Мощности", "B8_CAPACITY"),
]
dv_status = DataValidation(type="list", formula1='"не начат,в работе,готов"', allow_blank=True)
ws.add_data_validation(dv_status)
for i, (name, sheet) in enumerate(blocks):
    r = 28 + i
    ws.cell(row=r, column=1, value=name)
    ws.cell(row=r, column=2, value=sheet)
    c = ws.cell(row=r, column=3, value="в работе" if i == 0 else "не начат")
    inp(c)
    dv_status.add(c)
widths(ws, {"A": 30, "B": 110, "C": 14})

# ---------------- ANCHORS definition ----------------
# code, name, unit, block, aggregation, source, formula(template with {c} = column letter) or None
ANCH = [
    ("GDP_REAL", "Реальный ВВП", "млн грн, цены базового года", "B1/B2", "SUM", "Госстат", None),
    ("GDP_REAL_P", "Реальный ВВП по производству (сумма ВДС + чистые налоги)", "млн грн, цены базового года", "B1", "SUM", "Госстат", None),
    ("GDP_REAL_E", "Реальный ВВП по расходам (C + G + I + X − M)", "млн грн, цены базового года", "B2", "SUM", "Госстат", None),
    ("GDP_NOM", "Номинальный ВВП", "млн грн", "B3", "SUM", "Госстат", None),
    ("DEFL", "Дефлятор ВВП", "индекс, базовый год = 100", "B3", "AVG", "формула", "=IFERROR({c}{GDP_NOM}/{c}{GDP_REAL}*100,\"\")"),
    ("CPI_YOY", "ИПЦ", "% г/г, конец периода", "B3", "END", "Госстат", None),
    ("PPI_YOY", "Индекс цен производителей", "% г/г, конец периода", "B3", "END", "Госстат", None),
    ("FX_AVG", "Курс UAH/USD, средний", "грн за 1 USD", "B5", "AVG", "НБУ", None),
    ("FX_EOP", "Курс UAH/USD, конец периода", "грн за 1 USD", "B5", "END", "НБУ", None),
    ("NBU_RATE", "Учётная ставка НБУ", "%, конец периода", "B7", "END", "НБУ", None),
    ("LOAN_RATE", "Ставка по новым кредитам в гривне", "%, среднее", "B7", "AVG", "НБУ", None),
    ("DEP_RATE", "Ставка по новым депозитам в гривне", "%, среднее", "B7", "AVG", "НБУ", None),
    ("POP", "Население на подконтрольной территории", "млн человек", "B4", "AVG", "оценка", None),
    ("EMPL", "Занятые", "тыс. человек", "B4", "AVG", "Госстат / оценка", None),
    ("UNEMP", "Уровень безработицы", "%", "B4", "AVG", "Госстат / оценка", None),
    ("WAGE_NOM", "Средняя номинальная зарплата", "грн в месяц", "B4", "AVG", "Госстат", None),
    ("EXP_USD", "Экспорт товаров и услуг", "млн USD", "B5", "SUM", "НБУ (платёжный баланс)", None),
    ("IMP_USD", "Импорт товаров и услуг", "млн USD", "B5", "SUM", "НБУ (платёжный баланс)", None),
    ("TRADE_BAL", "Сальдо торговли товарами и услугами", "млн USD", "B5", "SUM", "формула", "=IF(COUNT({c}{EXP_USD},{c}{IMP_USD})<2,\"\",{c}{EXP_USD}-{c}{IMP_USD})"),
    ("RES_USD", "Международные резервы", "млн USD, конец периода", "B5", "END", "НБУ", None),
    ("GFCF", "Валовое накопление основного капитала", "млн грн", "B2", "SUM", "Госстат", None),
    ("CFC", "Потребление основного капитала (износ)", "млн грн", "B8", "SUM", "Госстат", None),
    ("BUDGET_DEF", "Дефицит госбюджета (+ = дефицит)", "млн грн", "B6", "SUM", "Минфин", None),
    ("FIN_NEED_USD", "Потребность во внешнем финансировании", "млн USD", "B6", "SUM", "расчёт", None),
    ("FIN_RECEIVED_USD", "Получено внешнего финансирования", "млн USD", "B6", "SUM", "Минфин / НБУ", None),
    ("FIN_GAP_USD", "Разрыв финансирования (потребность − получено)", "млн USD", "B6", "SUM", "формула", "=IF(COUNT({c}{FIN_NEED_USD},{c}{FIN_RECEIVED_USD})<2,\"\",{c}{FIN_NEED_USD}-{c}{FIN_RECEIVED_USD})"),
    ("DEBT_GDP", "Государственный и гарантированный долг", "% ВВП, конец периода", "B6", "END", "Минфин", None),
    ("ELEC_CONS", "Потребление электроэнергии", "млн кВт·ч", "B8", "SUM", "Укрэнерго / оценка", None),
    ("FREIGHT", "Перевезення вантажів", "млн т", "B8", "SUM", "Госстат (MacroStat)", None),
]
FIRST = 5
ROW = {a[0]: FIRST + i for i, a in enumerate(ANCH)}
LAST = FIRST + len(ANCH) - 1
QCOL0 = 8  # H

# ---------------- DICT ----------------
wd = wb.create_sheet("DICT")
wd["A1"] = "Словарь переменных"
wd["A1"].font = TITLE
wd["A2"] = "Одна строка на переменную. Новые переменные добавляются сюда до того, как появятся на других листах."
head(wd, 4, ["Код", "Название", "Единицы", "Блок", "Агрегация в год", "Тип", "Источник", "Ссылка / ряд", "Статус данных"])
dv_data = DataValidation(type="list", formula1='"нет,частично,загружено"', allow_blank=True)
wd.add_data_validation(dv_data)
for i, (code, name, unit, blk, agg, src, f) in enumerate(ANCH):
    r = 5 + i
    vals = [code, name, unit, blk, agg, "репер" + (" (формула)" if f else ""), src, None]
    for j, v in enumerate(vals):
        wd.cell(row=r, column=1 + j, value=v)
    c = wd.cell(row=r, column=9, value="—" if f else "нет")
    if not f:
        inp(c)
        dv_data.add(c)
wd.freeze_panes = "B5"
widths(wd, {"A": 18, "B": 52, "C": 28, "D": 8, "E": 16, "F": 18, "G": 24, "H": 40, "I": 14})

# ---------------- SCENARIOS ----------------
sc = wb.create_sheet("SCENARIOS")
sc["A1"] = "Сценарии и внешняя среда"
sc["A1"].font = TITLE
sc["A3"] = "Активный сценарий"
sc["C3"] = "База"
inp(sc["C3"])
sc["A4"] = "Последний фактический квартал"
sc["C4"] = "2026Q2"
inp(sc["C4"])
sc["D4"] = "Формат ГГГГQн. Всё позже — прогноз."
sc["A3"].font = B
sc["A4"].font = B

sc["A6"] = "Описание сценариев (черновое предложение)"
sc["A6"].font = B
sc["A6"].fill = SEC_FILL
head(sc, 7, ["Сценарий", "", "Война", "Энергетика и логистика", "Финансирование и люди"])
SC = [
    ("База",
     "Боевые действия продолжаются со снижающейся интенсивностью, режим прекращения огня с 2027 года без полного урегулирования",
     "Удары по энергетике сохраняются, но слабеют; дефицит закрывается импортом и распределённой генерацией; морской коридор работает с перебоями",
     "Внешнее финансирование по текущим обязательствам (МВФ, ЕС, ERA) с частичными задержками; эмиграция замедляется, возврат небольшой"),
    ("Пессимистичный",
     "Война высокой интенсивности до конца горизонта (2028)",
     "Систематические удары по генерации, сетям, портам и железной дороге; отключения для промышленности; морской экспорт сокращается",
     "Задержки и недобор внешнего финансирования, рост монетизации дефицита; продолжение оттока населения, рост мобилизации"),
    ("Оптимистичный",
     "Прекращение огня в начале 2027 года с гарантиями безопасности",
     "Удары прекращаются, восстановление генерации и портов, снижение логистических издержек",
     "Финансирование восстановления и рост прямых инвестиций; часть мигрантов возвращается; оборонные расходы постепенно снижаются"),
]
for i, (n, w, e, f) in enumerate(SC):
    r = 8 + i
    sc.cell(row=r, column=1, value=n).font = B
    for j, v in enumerate([w, e, f]):
        c = sc.cell(row=r, column=3 + j, value=v)
        c.alignment = WRAP
        inp(c)
    sc.row_dimensions[r].height = 62
dv_sc = DataValidation(type="list", formula1="=$A$8:$A$10", allow_blank=False)
sc.add_data_validation(dv_sc)
dv_sc.add("C3")

sc["A12"] = "Сценарные допущения по годам"
sc["A12"].font = B
sc["A12"].fill = SEC_FILL
sc["A13"] = "Значения заполняются после загрузки истории и калибровки. Колонки «Активный» выбирают значения сценария из C3."
head(sc, 14, ["Параметр", "Единицы", "База", "", "", "Пессимистичный", "", "", "Оптимистичный", "", "", "Активный", "", ""])
head(sc, 15, ["", ""] + [2026, 2027, 2028] * 4)
PARAMS = [
    ("Квартал прекращения активных боевых действий", "ГГГГQн или «нет»"),
    ("Интенсивность ударов по энергетике", "индекс 0–1"),
    ("Доступная мощность энергосистемы в пик", "ГВт"),
    ("Пропускная способность морских портов", "млн т в год"),
    ("Обещанное внешнее финансирование бюджета", "млн USD"),
    ("Расходы на оборону", "% ВВП"),
    ("Чистая миграция", "тыс. человек в год"),
    ("Мобилизация (чистый отток из рабочей силы)", "тыс. человек в год"),
    ("Урожай зерновых и масличных", "млн т"),
    ("Цена пшеницы", "USD/т"),
    ("Цена стали (заготовка)", "USD/т"),
    ("Цена газа TTF", "EUR/МВт·ч"),
    ("Рост ВВП еврозоны", "%"),
    ("Инфляция еврозоны", "%"),
    ("Ставка ФРС", "%, конец года"),
]
for i, (p, u) in enumerate(PARAMS):
    r = 16 + i
    sc.cell(row=r, column=1, value=p)
    sc.cell(row=r, column=2, value=u)
    for col in range(3, 12):
        inp(sc.cell(row=r, column=col))
    for k in range(3):
        b, ps, o = L(3 + k), L(6 + k), L(9 + k)
        sc.cell(row=r, column=12 + k,
                value=f'=IFERROR(CHOOSE(MATCH($C$3,$A$8:$A$10,0),{b}{r},{ps}{r},{o}{r}),"")')
sc.freeze_panes = "C16"
widths(sc, {"A": 46, "B": 20, "C": 34, "D": 34, "E": 34})
for col in range(6, 15):
    sc.column_dimensions[L(col)].width = 12

# ---------------- ANCHORS ----------------
wa = wb.create_sheet("ANCHORS")
wa["A1"] = "Реперные переменные (квартал)"
wa["A1"].font = TITLE
wa["A2"] = "Единственный канал связи между блоками. Факт — ввод из источников; прогноз — пишет лист блока."
wa.cell(row=3, column=7, value="Фаза").font = B
wa.cell(row=2, column=7, value="Год").font = B
head(wa, 4, ["Код", "Название", "Единицы", "Блок", "Агрегация", "Источник", "Тип"] + QUARTERS)
for k, q in enumerate(QUARTERS):
    col = L(QCOL0 + k)
    wa.cell(row=3, column=QCOL0 + k, value=f'=IF({col}$4<=SCENARIOS!$C$4,"факт","прогноз")')
    wa.cell(row=2, column=QCOL0 + k, value=f'=VALUE(LEFT({col}$4,4))')
for i, (code, name, unit, blk, agg, src, f) in enumerate(ANCH):
    r = FIRST + i
    for j, v in enumerate([code, name, unit, blk, agg, src, "формула" if f else "данные"]):
        wa.cell(row=r, column=1 + j, value=v)
    for k in range(len(QUARTERS)):
        col = L(QCOL0 + k)
        c = wa.cell(row=r, column=QCOL0 + k)
        if f:
            c.value = f.format(c=col, **{kk: vv for kk, vv in ROW.items()})
qend = L(QCOL0 + len(QUARTERS) - 1)
from openpyxl.formatting.rule import FormulaRule
wa.conditional_formatting.add(f"H{FIRST}:{qend}{LAST}", FormulaRule(formula=[f'$G{FIRST}="данные"'], fill=INPUT_FILL, font=INPUT_FONT))
wa.freeze_panes = "H5"
widths(wa, {"A": 18, "B": 44, "C": 24, "D": 7, "E": 10, "F": 18, "G": 9})

# ---------------- ANNUAL ----------------
wy = wb.create_sheet("ANNUAL")
wy["A1"] = "Годовая агрегация реперов"
wy["A1"].font = TITLE
wy["A2"] = "SUM — сумма кварталов, AVG — среднее, END — значение IV квартала. Пусто, если года нет целиком."
head(wy, 4, ["Код", "Название", "Единицы", "Агрегация"] + YEARS)
for i in range(len(ANCH)):
    r = FIRST + i
    wy.cell(row=r, column=1, value=f"=ANCHORS!A{r}")
    wy.cell(row=r, column=2, value=f"=ANCHORS!B{r}")
    wy.cell(row=r, column=3, value=f"=ANCHORS!C{r}")
    wy.cell(row=r, column=4, value=f"=ANCHORS!E{r}")
    for yi, y in enumerate(YEARS):
        yc = L(5 + yi)
        row = f"ANCHORS!$H{r}:$BK{r}"
        yrs = "ANCHORS!$H$2:$BK$2"
        wy.cell(row=r, column=5 + yi, value=(
            f'=IF(SUMPRODUCT(({yrs}={yc}$4)*ISNUMBER({row}))<4,"",'
            f'IF($D{r}="SUM",SUMIFS({row},{yrs},{yc}$4),'
            f'IF($D{r}="AVG",AVERAGEIFS({row},{yrs},{yc}$4),'
            f'INDEX({row},MATCH({yc}$4&"Q4",ANCHORS!$H$4:$BK$4,0)))))'))
wy.freeze_panes = "E5"
widths(wy, {"A": 18, "B": 44, "C": 24, "D": 10})

# ---------------- CHECKS ----------------
wc = wb.create_sheet("CHECKS")
wc["A1"] = "Проверки согласованности и мощностей (год)"
wc["A1"].font = TITLE
wc["A2"] = "Порог — ввод (черновое значение, откалибровать на истории). Статус: OK / ВНИМАНИЕ / нет данных."
head(wc, 4, ["Проверка", "Как считается", "Порог"] + YEARS)
YC0 = 4  # D


def a(code, yi, lag=0):
    col = L(5 + yi - lag)  # ANNUAL year column
    return f'INDEX(ANNUAL!{col}${FIRST}:{col}${LAST},MATCH("{code}",ANNUAL!$A${FIRST}:$A${LAST},0))'


CHECKS = [
    ("Сверху вниз / снизу вверх", None, None, None, None),
    ("Реальный ВВП: производство против расходов", "ВВП по производству / ВВП по расходам − 1", 0.01,
     lambda yi: f'=IFERROR({a("GDP_REAL_P", yi)}/{a("GDP_REAL_E", yi)}-1,"")', "abs"),
    ("Мощности", None, None, None, None),
    ("Люди: рост производительности труда", "Рост (ВВП / занятые) г/г; не выше порога", 0.06,
     lambda yi: None if yi == 0 else f'=IFERROR(({a("GDP_REAL", yi)}/{a("EMPL", yi)})/({a("GDP_REAL", yi, 1)}/{a("EMPL", yi, 1)})-1,"")', "max"),
    ("Энергия: изменение энергоёмкости", "Изменение (электроэнергия / ВВП) г/г; не ниже порога", -0.08,
     lambda yi: None if yi == 0 else f'=IFERROR(({a("ELEC_CONS", yi)}/{a("GDP_REAL", yi)})/({a("ELEC_CONS", yi, 1)}/{a("GDP_REAL", yi, 1)})-1,"")', "min"),
    ("Логистика: изменение транспортоёмкости", "Изменение (грузооборот / ВВП) г/г; не ниже порога", -0.08,
     lambda yi: None if yi == 0 else f'=IFERROR(({a("FREIGHT", yi)}/{a("GDP_REAL", yi)})/({a("FREIGHT", yi, 1)}/{a("GDP_REAL", yi, 1)})-1,"")', "min"),
    ("Капитал: чистые инвестиции", "(ВНОК − износ) / номинальный ВВП; не ниже порога", 0.0,
     lambda yi: f'=IFERROR(({a("GFCF", yi)}-{a("CFC", yi)})/{a("GDP_NOM", yi)},"")', "min"),
    ("Финансирование", None, None, None, None),
    ("Разрыв внешнего финансирования", "Потребность − получено, млн USD; не выше порога", 0,
     lambda yi: f'=IFERROR({a("FIN_GAP_USD", yi)},"")', "max"),
]
r = 5
for name, how, thr, fn, kind in CHECKS:
    if fn is None:
        c = wc.cell(row=r, column=1, value=name)
        c.font = B
        for col in range(1, 4 + len(YEARS)):
            wc.cell(row=r, column=col).fill = SEC_FILL
        r += 1
        continue
    wc.cell(row=r, column=1, value=name)
    wc.cell(row=r, column=2, value=how)
    t = wc.cell(row=r, column=3, value=thr)
    inp(t)
    wc.cell(row=r + 1, column=1, value="  статус")
    for yi in range(len(YEARS)):
        col = L(YC0 + yi)
        v = fn(yi)
        if v:
            wc.cell(row=r, column=YC0 + yi, value=v)
            cond = {"abs": f"ABS({col}{r})<=$C${r}", "max": f"{col}{r}<=$C${r}", "min": f"{col}{r}>=$C${r}"}[kind]
            wc.cell(row=r + 1, column=YC0 + yi, value=f'=IF({col}{r}="","нет данных",IF({cond},"OK","ВНИМАНИЕ"))')
    r += 2
end_col = L(YC0 + len(YEARS) - 1)
wc.conditional_formatting.add(f"D5:{end_col}{r}", CellIsRule(operator="equal", formula=['"ВНИМАНИЕ"'], fill=PatternFill("solid", fgColor="F4CCCC")))
wc.conditional_formatting.add(f"D5:{end_col}{r}", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="D9EAD3")))
wc.freeze_panes = "D5"
widths(wc, {"A": 44, "B": 52, "C": 9})

# ---------------- COEFS ----------------
wk = wb.create_sheet("COEFS")
wk["A1"] = "Коэффициенты уравнений"
wk["A1"].font = TITLE
wk["A2"] = "Заполняет Python после оценки. Одна строка — один коэффициент. Листы блоков берут значения по ID уравнения и регрессору."
head(wk, 4, ["ID уравнения", "Блок", "Зависимая", "Регрессор", "Коэффициент", "Ст. ошибка", "t-стат", "Период оценки", "Метод", "R²", "Дата оценки", "Комментарий"])
wk.freeze_panes = "A5"
widths(wk, {"A": 14, "B": 8, "C": 16, "D": 18, "H": 16, "I": 12, "K": 12, "L": 40})

# ---------------- Block sheets ----------------
BLK = [
    ("B1_REAL", "1. Реальный сектор", "ВДС по отраслям: сельское хозяйство, добыча, переработка, энергетика, строительство, транспорт, торговля",
     "SCENARIOS, EMPL, GFCF, ELEC_CONS, FREIGHT", "GDP_REAL_P"),
    ("B2_DEMAND", "2. Спрос", "Потребление домохозяйств (розница), госпотребление, ВНОК, экспорт и импорт в реальном выражении",
     "WAGE_NOM, CPI_YOY, LOAN_RATE, BUDGET_DEF, FX_AVG", "GDP_REAL_E, GFCF"),
    ("B3_PRICES", "3. Цены", "ИПЦ (базовый, продукты, регулируемые), ИЦП, дефлятор, номинальный ВВП",
     "FX_AVG, WAGE_NOM, NBU_RATE, SCENARIOS (цены сырья)", "CPI_YOY, PPI_YOY, GDP_NOM"),
    ("B4_LABOUR", "4. Рынок труда", "Население, занятость, безработица, зарплаты",
     "SCENARIOS (миграция, мобилизация), GDP_REAL, CPI_YOY", "POP, EMPL, UNEMP, WAGE_NOM"),
    ("B5_EXTERNAL", "5. Внешний сектор", "Экспорт и импорт по товарным группам, сальдо, резервы, курс",
     "SCENARIOS (урожай, цены, порты), GDP_REAL_E, FIN_RECEIVED_USD", "EXP_USD, IMP_USD, RES_USD, FX_AVG, FX_EOP"),
    ("B6_FISCAL", "6. Бюджет и долг", "Доходы, расходы (оборона), дефицит, потребность и получение финансирования, долг",
     "GDP_NOM, WAGE_NOM, IMP_USD, FX_AVG, NBU_RATE, SCENARIOS", "BUDGET_DEF, FIN_NEED_USD, FIN_RECEIVED_USD, DEBT_GDP"),
    ("B7_MONETARY", "7. Монетарный блок", "Учётная ставка (правило типа Тейлора), ставки кредитов и депозитов, ОВДП, ликвидность, M2",
     "CPI_YOY, GDP_REAL, BUDGET_DEF, RES_USD", "NBU_RATE, LOAN_RATE, DEP_RATE"),
    ("B8_CAPACITY", "8. Мощности", "Потолок выпуска по людям, энергии, логистике и капиталу (минимум из четырёх)",
     "EMPL, ELEC_CONS, FREIGHT, GFCF, CFC, SCENARIOS", "потолок выпуска по отраслям → CHECKS"),
]
for sheet, name, what, ins, outs in BLK:
    w = wb.create_sheet(sheet)
    w["A1"] = name
    w["A1"].font = TITLE
    for rr, (k, v) in enumerate([("Что считает", what), ("Берёт из ANCHORS / SCENARIOS", ins), ("Пишет в ANCHORS", outs), ("Статус", "не начат"), ("", "Уравнения и разложение по вкладам появятся здесь")], start=3):
        w.cell(row=rr, column=1, value=k).font = B
        w.cell(row=rr, column=2, value=v)
    widths(w, {"A": 30, "B": 100})

# ---------------- SOURCES ----------------
so = wb.create_sheet("SOURCES")
so["A1"] = "Источники данных"
so["A1"].font = TITLE
head(so, 3, ["Источник", "Что берём", "Ссылка", "Способ загрузки"])
SRC = [
    ("НБУ, API", "Курсы, ставки, денежные агрегаты", "https://bank.gov.ua/ua/open-data/api-dev", "Python, API"),
    ("НБУ", "Интервенции", "https://bank.gov.ua/files/Finmarket/InterventionsResults.xlsx", "Python, xlsx"),
    ("НБУ", "Операции по регулированию ликвидности", "https://bank.gov.ua/files/PInterestRates_Facts.xls", "Python, xls"),
    ("НБУ", "Остатки на корсчетах", "https://bank.gov.ua/files/Vzal.xlsx", "Python, xlsx"),
    ("НБУ", "Ставки по кредитам и депозитам", "https://bank.gov.ua/files/4-Financial_markets.xlsx", "Python, xlsx"),
    ("НБУ", "ОВДП в обращении", "https://bank.gov.ua/files/Depositary/T-bills_debt.xlsx", "Python, xlsx"),
    ("НБУ", "Денежные агрегаты M0–M3", "https://bank.gov.ua/files/SDDS/AA_BS.xlsx", "Python, xlsx"),
    ("НБУ", "Показатели валютного рынка", "https://bank.gov.ua/files/rates_final.xlsx", "Python, xlsx"),
    ("Госстат", "ВВП, ВДС, ИПЦ, ИЦП, занятость, зарплаты, отрасли", "https://www.ukrstat.gov.ua/", "Python, xlsx (адрес проверить)"),
    ("Минфин", "Бюджет, долг, финансирование", "https://mof.gov.ua/", "Python / вручную"),
    ("Минфин (index.minfin.com.ua)", "ОВДП, погашения", "https://index.minfin.com.ua/ua/finance/bonds/", "Python"),
    ("МВФ", "WEO, отчёты программы EFF, мировые допущения", "https://www.imf.org/en/Publications/WEO", "Python, API"),
    ("Всемирный банк", "WDI, оценки ущерба (RDNA)", "https://api.worldbank.org/v2/", "Python, API"),
    ("Пользователь", "MacroStat (накопленные знания)", "https://docs.google.com/spreadsheets/d/11d11t-sQ7ESeZWnu3KkjWRjWEVuWwaajzuKniL39F_k/edit", "вручную"),
]
for i, row in enumerate(SRC):
    for j, v in enumerate(row):
        so.cell(row=4 + i, column=1 + j, value=v)
widths(so, {"A": 28, "B": 48, "C": 70, "D": 26})

for w in wb.worksheets:
    w.sheet_view.showGridLines = True
wb.save(OUT)
print(OUT)
