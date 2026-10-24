"""NBU statistical xlsx files that carry history the API lacks.

- ES/Trade_m.xlsx: monthly goods exports and imports by commodity group, USD million,
  from 2010 (BPM6). The API has only goods and services totals.
- macro/GDP_y.xlsx: annual GDP (Gosstat) back to 2001: nominal, real index, deflator.
  The API's quarterly GDP starts in 2010.
- macro/PPI_q.xlsx: quarterly producer prices 2007-2012; the API starts in 2013.
- macro/Wages_m.xlsx, Wages_y.xlsx: average wage monthly from 2010 and by year from
  1995; the API starts in 2013.
- macro/Number_m.xlsx: average number of staff employees from 2010, through the war years.
- macro/CCPI_m.xlsx: core CPI from 2008; the API's core series has gaps (2019-06 reads -92.6).
- macro/Surveys_price.xlsx, Surveys_rate.xlsx: 12-month-ahead inflation and UAH/USD
  expectations of banks, businesses, households and financial analysts, from 2013.

The files are replaced in place at the same address, so each is downloaded to the cache
and read by row labels, not positions.
"""
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import requests

BASE = "https://bank.gov.ua/files"
FILES = {"trade_m": "ES/Trade_m.xlsx", "gdp_y": "macro/GDP_y.xlsx", "ppi_q": "macro/PPI_q.xlsx",
         "wages_m": "macro/Wages_m.xlsx", "wages_y": "macro/Wages_y.xlsx", "staff_m": "macro/Number_m.xlsx",
         "ccpi_m": "macro/CCPI_m.xlsx", "surveys_price": "macro/Surveys_price.xlsx", "surveys_rate": "macro/Surveys_rate.xlsx"}

MONTHS_UA = ["січ", "лют", "бер", "кві", "тра", "чер", "лип", "сер", "вер", "жов", "лис", "гру"]  # first three letters

# English row labels in Trade_m sheets 1.1 (exports) and 1.2 (imports)
TRADE_GROUPS = {
    "TOTAL": "TOTAL",
    "AGRO": "Agricultural products",
    "MINERAL": "Mineral products",
    "CHEM": "Chemicals",
    "WOOD": "Timber and wood products",
    "INDUSTRIAL": "Industrial goods",
    "METALS": "Ferrrous and nonferrous metals",
    "MACHINERY": "Machinery and equipment",
    "OTHER": "Other",
}


def path(cache: Path, name: str, refresh: bool) -> Path:
    p = cache / "nbu_files" / Path(FILES[name]).name
    if refresh or not p.exists():
        p.parent.mkdir(parents=True, exist_ok=True)
        r = requests.get(f"{BASE}/{FILES[name]}", timeout=120)
        r.raise_for_status()
        p.write_bytes(r.content)
    return p


def _rows(p: Path, sheet: str) -> list[tuple]:
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return rows


def _month(v) -> int | None:
    if not isinstance(v, str):
        return None
    v = v.strip().lower().rstrip(".*")
    if "-" in v or len(v) < 3 or not v.isalpha():  # 'січ-серп' is a year-to-date column
        return None
    return MONTHS_UA.index(v[:3]) + 1 if v[:3] in MONTHS_UA else None


def trade_groups_m(p: Path, sheet: str) -> pd.DataFrame:
    """Monthly goods trade by commodity group, USD million. sheet '1.1' exports, '1.2' imports.

    Years label the first month column of their block (row 5, a number or a text), months
    the row below (spelled in full or abbreviated); the running year carries over until the
    next label. Summary columns on the right have no
    month name and are skipped."""
    rows = _rows(p, sheet)
    hi = next(i for i, r in enumerate(rows) if any(_month(v) for v in r))
    years, months = rows[hi - 1], rows[hi]
    cols, year = {}, None
    for j, mv in enumerate(months):
        y = years[j] if j < len(years) else None
        if isinstance(y, str) and y.strip().isdigit():
            y = int(y.strip())
        m = _month(mv)
        if isinstance(y, (int, float)) and 2000 < y < 2100 and m:
            year = int(y)
        if m and year:
            cols[j] = pd.Period(f"{year}-{m:02d}", "M")
    out = {}
    for key, label in TRADE_GROUPS.items():
        r = next(r for r in rows[hi + 1:] if isinstance(r[2], str) and r[2].strip().startswith(label))
        out[key] = {per: float(r[j]) for j, per in cols.items() if isinstance(r[j], (int, float))}
    return pd.DataFrame(out).sort_index()


def gdp_y(p: Path) -> pd.DataFrame:
    """Annual GDP, SNA 2008 rows: nominal (UAH million), real index and deflator (% to previous year)."""
    rows = _rows(p, next(s for s in openpyxl.load_workbook(p, read_only=True).sheetnames if s.startswith("GDP")))
    head = next(r for r in rows if sum(isinstance(v, (int, float)) and 1990 < v < 2100 for v in r) > 5)
    sna = [r for r in rows if isinstance(r[0], str) and "СНР 2008" in r[0]]
    out = {}
    for key, r in zip(["GDP_NOM", "GDP_REAL_IDX", "DEFL_IDX"], sna):
        out[key] = {int(h): float(v) for h, v in zip(head, r) if isinstance(h, (int, float)) and isinstance(v, (int, float))}
    return pd.DataFrame(out).sort_index()


def ppi_q(p: Path) -> pd.DataFrame:
    """Quarterly PPI 2007-2012, % change to the same quarter of the previous year.

    Sheet 2 holds y/y changes by branch; returned columns: TOTAL (industry), MINING,
    MANUF, METALS (metallurgy and metal products) where the row exists."""
    rows = _rows(p, "2")
    hi = next(i for i, r in enumerate(rows) if any(isinstance(v, str) and re.fullmatch(r"[IV]+\.\d{4}", v.strip()) for v in r))
    roman = {"I": 1, "II": 2, "III": 3, "IV": 4}
    cols = {}
    for j, v in enumerate(rows[hi]):
        if isinstance(v, str) and re.fullmatch(r"[IV]+\.\d{4}", v.strip()):
            q, y = v.strip().split(".")
            cols[j] = f"{y}Q{roman[q]}"
    labels = {"TOTAL": "Промисловість", "MINING": "Добувна промисловість", "MANUF": "Переробна промисловість",
              "METALS": "Металургійне виробництво"}
    out = {}
    for key, lab in labels.items():
        r = next((r for r in rows[hi + 1:] if isinstance(r[0], str) and r[0].strip().startswith(lab)), None)
        if r is not None:
            out[key] = {per: float(r[j]) for j, per in cols.items() if isinstance(r[j], (int, float))}
    return pd.DataFrame(out).sort_index()


def _monthly_row(p: Path, label: str, sheets: list[str] | None = None) -> pd.Series:
    """First row starting with `label` under a header of month dates, over the given
    sheets (default all but the contents sheet). Sheets run newest first; the first sheet
    that has a month wins."""
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = {}
    for name in sheets or wb.sheetnames[1:]:
        rows = list(wb[name].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows) if sum(isinstance(v, datetime) for v in r) >= 3), None)
        if hi is None:
            continue
        cols = {j: pd.Period(v, "M") for j, v in enumerate(rows[hi]) if isinstance(v, datetime)}
        r = next((r for r in rows[hi + 1:] if isinstance(r[0], str) and r[0].strip().startswith(label)), None)
        if r is None:
            continue
        for j, per in cols.items():
            if j < len(r) and isinstance(r[j], (int, float)) and per not in out:
                out[per] = float(r[j])
    wb.close()
    return pd.Series(out, dtype=float).sort_index()


def wages_m(p: Path) -> pd.Series:
    """Average monthly wage, UAH, whole economy, from 2010 (KVED 2005 and 2010 sheets)."""
    return _monthly_row(p, "Середня заробітна плата")


def staff_m(p: Path) -> pd.Series:
    """Average registered number of staff employees, thousand, from 2010 (sheets 1-2 hold
    levels, 3-4 indices). Gosstat's enterprise survey: continues after the labour force
    survey stopped in 2021, so it is the employment signal for the war years."""
    return _monthly_row(p, "Середньооблікова", ["1", "2"])


def wages_y(p: Path) -> pd.Series:
    """Average monthly wage by year, UAH: the KVED 2010 row where it has a value, else the
    KVED 2005 row above it (years before 2010)."""
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    rows = list(wb[next(s for s in wb.sheetnames if s.startswith("Wages"))].iter_rows(values_only=True))
    wb.close()
    hi = next(i for i, r in enumerate(rows) if sum(isinstance(v, (datetime, str)) and re.match(r"(19|20)\d\d", str(v)) is not None for v in r[1:]) >= 5)
    years = {j: int(str(v)[:4]) for j, v in enumerate(rows[hi]) if j and re.match(r"(19|20)\d\d", str(v))}
    out = {}
    for r in rows[hi + 1:hi + 3]:  # KVED 2005, then KVED 2010
        for j, y in years.items():
            if j < len(r) and isinstance(r[j], (int, float)):
                out[y] = float(r[j])
    return pd.Series(out, dtype=float).sort_index()


def core_cpi_yoy_m(p: Path) -> pd.Series:
    """Core CPI, % to the same month of the previous year (published from 2009)."""
    return _monthly_row(p, "до відповідного місяця попереднього року", ["1"])


# column headers of the Surveys_* sheet "UKR"
SURVEY_GROUPS = {"BANKS": "банків", "BUSINESS": "підприємств", "HH": "домогосподарств", "ANALYSTS": "фінансових аналітиків"}


def surveys_m(p: Path) -> pd.DataFrame:
    """Expectations 12 months ahead by respondent group, one column per group, monthly.
    Each group is surveyed at its own frequency (households monthly, banks and businesses
    quarterly, analysts in most months); months without a survey hold '#N/A' and stay empty."""
    rows = _rows(p, "UKR")
    hi = next(i for i, r in enumerate(rows) if any(isinstance(v, str) and v.strip() == "банків" for v in r))
    cols = {k: next(j for j, v in enumerate(rows[hi]) if isinstance(v, str) and v.strip() == lab) for k, lab in SURVEY_GROUPS.items()}
    out = {}
    for r in rows[hi + 1:]:
        if not isinstance(r[0], datetime):
            continue
        out[pd.Period(r[0], "M")] = {k: float(r[j]) if isinstance(r[j], (int, float)) else None for k, j in cols.items()}
    return pd.DataFrame.from_dict(out, orient="index", dtype=float).sort_index()
