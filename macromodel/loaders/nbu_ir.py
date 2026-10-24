"""CPI components by the NBU's own split, from the data files of the Inflation Report.

Each quarterly report has a "Tables_Charts_YYYY-Qn.xlsx" with the data behind its charts.
The NBU splits CPI into core, raw food, administered prices and fuel (and, since 2023,
also shows processed food, nonfood goods and services). No statistical file carries
this split, so the history is pieced together from the reports:
- quarterly: the forecast charts ("<component>, %" with a "Річна зміна" column) give
  end-of-quarter % y/y from 2017Q2 (fuel from 2019Q2). Rows before the report's quarter
  are actual, later rows are forecast and are dropped. A later report wins.
- monthly: the "Selected CPI components, % yoy" chart, from 2023.
Sheet numbers move between reports, so sheets are found by their captions.
"""
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import requests

URL = "https://bank.gov.ua/admin_uploads/article/Tables_Charts_{y}-Q{q}.xlsx"
FIRST_YEAR = 2019  # earlier reports have no data file

# caption text of the quarterly forecast charts -> code
Q_CAPTIONS = {"CORE": r"Базовий ІСЦ, %", "RAW_FOOD": r"сирі продовольчі товари", "ADMIN": r"Адміністративно регульовані ціни",
              "FUEL": r"Ціни на (паливо|пальне)"}
# English legend of the monthly "Selected CPI components" chart -> code
M_LABELS = {"Processed foods": "PROC_FOOD", "Raw foods": "RAW_FOOD", "Administered prices": "ADMIN", "Nonfood products": "NONFOOD",
            "Nonfood prices": "NONFOOD", "Services": "SERVICES", "Service prices": "SERVICES", "Fuel prices": "FUEL", "Food prices": "FOOD"}
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4}


def download(cache: Path, refresh: bool, today: datetime | None = None) -> list[Path]:
    """Every report's data file, oldest first; reports without one are skipped."""
    today = today or datetime.today()
    d = cache / "nbu_ir"
    d.mkdir(parents=True, exist_ok=True)
    out = []
    for y in range(FIRST_YEAR, today.year + 1):
        for q in range(1, 5):
            if (y, q) > (today.year, (today.month - 1) // 3 + 1):
                break
            p = d / f"Tables_Charts_{y}-Q{q}.xlsx"
            if not p.exists() or (refresh and y == today.year):
                r = requests.get(URL.format(y=y, q=q), timeout=120)
                if r.status_code != 200:
                    continue
                p.write_bytes(r.content)
            out.append(p)
    return out


def report_quarter(p: Path) -> pd.Period:
    y, q = re.search(r"(\d{4})-Q(\d)", p.name).groups()
    return pd.Period(f"{y}Q{q}", "Q")


def _quarter(label: str) -> pd.Period | None:
    m = re.fullmatch(r"(I{1,3}|IV)\.(\d{2})", label.strip())
    return pd.Period(f"20{m.group(2)}Q{ROMAN[m.group(1)]}", "Q") if m else None


def _parse(p: Path) -> tuple[dict, dict]:
    """Quarterly and monthly values of one report, {code: {period: value}}. Loading a report
    takes seconds, so the result is kept next to it and reused while the file is unchanged."""
    cached = p.with_suffix(".parsed.pkl")
    if cached.exists() and cached.stat().st_mtime >= p.stat().st_mtime:
        return pd.read_pickle(cached)
    last_actual = report_quarter(p) - 1
    q: dict[str, dict] = {}
    m: dict[str, dict] = {}
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    for ws in wb:
        top = list(ws.iter_rows(max_row=3, values_only=True))
        if not top:
            continue
        if "Річна зміна" in [str(v).strip() for v in top[0] if v is not None]:
            head = " ".join(str(v) for v in top[0] if isinstance(v, str))
            code = next((k for k, pat in Q_CAPTIONS.items() if re.search(pat, head)), None)
            if code is None:
                continue
            rows = list(ws.iter_rows(values_only=True))
            col = [str(v).strip() if v is not None else "" for v in rows[0]].index("Річна зміна")
            first = next(((i, _quarter(r[0])) for i, r in enumerate(rows) if isinstance(r[0], str) and _quarter(r[0])), None)
            if first is None:
                continue
            i0, q0 = first
            for i in range(i0, len(rows)):
                per = q0 + (i - i0)
                if per > last_actual:
                    break
                v = rows[i][col] if col < len(rows[i]) else None
                if isinstance(v, (int, float)):
                    q.setdefault(code, {})[per] = float(v)
        elif len(top) == 3 and any(isinstance(v, str) and "Selected CPI components" in v for v in top[2]):
            en = [str(v).strip().rstrip("*").replace(" (RHS)", "") if isinstance(v, str) else "" for v in top[2]]
            cols = {j: M_LABELS[x] for j, x in enumerate(en) if x in M_LABELS}
            for r in ws.iter_rows(min_row=4, values_only=True):
                if not r or not isinstance(r[0], datetime):
                    continue
                per = pd.Period(r[0], "M")
                for j, code in cols.items():
                    if j < len(r) and isinstance(r[j], (int, float)):
                        m.setdefault(code, {})[per] = float(r[j])
    wb.close()
    pd.to_pickle((q, m), cached)
    return q, m


def _merge(parts: list[dict]) -> pd.DataFrame:
    out: dict[str, dict] = {}
    for part in parts:  # oldest report first, so later reports overwrite
        for code, vals in part.items():
            out.setdefault(code, {}).update(vals)
    return pd.DataFrame(out).sort_index()


def quarterly(paths: list[Path]) -> pd.DataFrame:
    """End-of-quarter % y/y by component, actual quarters only; later reports win."""
    return _merge([_parse(p)[0] for p in paths])


def monthly(paths: list[Path]) -> pd.DataFrame:
    """Monthly % y/y from the "Selected CPI components" chart; later reports win."""
    return _merge([_parse(p)[1] for p in paths])
